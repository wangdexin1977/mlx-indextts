from __future__ import annotations

import shutil
import subprocess
from pathlib import Path

import pytest

from mlx_indextts.document_import import (
    DocumentImportError,
    count_chinese_characters,
    count_effective_characters,
    import_document,
    select_document_text,
)


def test_txt_gbk_and_character_counts(tmp_path: Path):
    source = tmp_path / "sample.txt"
    source.write_bytes("第一章\n今天天气很好。\n\n第二段。".encode("gb18030"))

    result = import_document(source)

    assert result.file_type == "TXT"
    assert "今天天气很好" in result.text
    assert result.character_count == count_effective_characters(result.text)
    assert result.chinese_character_count == count_chinese_characters(result.text)


def test_markdown_removes_markup_and_code(tmp_path: Path):
    source = tmp_path / "sample.md"
    source.write_text(
        "# 章节标题\n\n**重要内容**，[链接文字](https://example.com)。\n\n"
        "```python\nprint('skip me')\n```\n",
        encoding="utf-8",
    )

    result = import_document(source)

    assert "章节标题" in result.text
    assert "重要内容" in result.text
    assert "链接文字" in result.text
    assert "https://" not in result.text
    assert "skip me" not in result.text
    assert "**" not in result.text


def _create_docx(path: Path) -> None:
    from docx import Document

    document = Document()
    document.add_heading("第一章", level=1)
    document.add_paragraph("这是第一章的正文。")
    document.add_heading("第二章", level=1)
    document.add_paragraph("这是第二章的正文。")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "表格字段"
    table.cell(0, 1).text = "表格内容"
    document.save(path)


def test_docx_chapters_and_tables(tmp_path: Path):
    source = tmp_path / "sample.docx"
    _create_docx(source)

    result = import_document(source)

    assert [chapter.title for chapter in result.chapters] == ["第一章", "第二章"]
    assert "表格字段，表格内容" in result.text
    selected = select_document_text(result, ["第二章"])
    assert "第二章" in selected
    assert "第一章的正文" not in selected


def test_legacy_doc_conversion(tmp_path: Path):
    source_docx = tmp_path / "legacy-source.docx"
    _create_docx(source_docx)
    soffice = shutil.which("soffice")
    if not soffice:
        pytest.skip("LibreOffice is unavailable")
    profile = tmp_path / "lo-profile"
    process = subprocess.run(
        [
            soffice,
            f"-env:UserInstallation=file://{profile}",
            "--headless",
            "--convert-to",
            "doc:MS Word 97",
            "--outdir",
            str(tmp_path),
            str(source_docx),
        ],
        capture_output=True,
        text=True,
        timeout=120,
    )
    assert process.returncode == 0, process.stderr
    source_doc = tmp_path / "legacy-source.doc"
    assert source_doc.exists()

    result = import_document(source_doc)

    assert result.file_type == "Word DOC"
    assert "第一章" in result.text
    assert "第二章" in result.text


def test_text_pdf_without_ocr(tmp_path: Path):
    import pymupdf

    source = tmp_path / "text.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "This is a searchable PDF page with enough readable text.")
    document.save(source)
    document.close()

    result = import_document(source)

    assert result.file_type == "PDF"
    assert not result.used_ocr
    assert "searchable PDF" in result.text


def test_scanned_pdf_uses_local_chinese_ocr(tmp_path: Path):
    from PIL import Image, ImageDraw, ImageFont

    font_candidates = [
        "/System/Library/Fonts/PingFang.ttc",
        "/System/Library/Fonts/STHeiti Medium.ttc",
        "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
    ]
    font_path = next((path for path in font_candidates if Path(path).exists()), None)
    if not font_path or not shutil.which("tesseract"):
        pytest.skip("Chinese font or Tesseract is unavailable")

    image = Image.new("RGB", (1654, 2339), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(font_path, 64)
    draw.text((140, 260), "扫描文档本机识别测试", fill="black", font=font)
    draw.text((140, 390), "这是一段需要转换成语音的中文内容。", fill="black", font=font)
    source = tmp_path / "scan.pdf"
    image.save(source, "PDF", resolution=150.0)

    result = import_document(source)

    assert result.used_ocr
    assert result.ocr_pages == [1]
    assert "本机" in result.text
    assert "中文" in result.text


def _create_epub(path: Path) -> None:
    from ebooklib import epub

    book = epub.EpubBook()
    book.set_identifier("indextts-test")
    book.set_title("电子书测试")
    book.set_language("zh-CN")
    chapter_one = epub.EpubHtml(title="第一章", file_name="one.xhtml", lang="zh-CN")
    chapter_one.content = "<h1>第一章</h1><p>第一章正文内容。</p>"
    chapter_two = epub.EpubHtml(title="第二章", file_name="two.xhtml", lang="zh-CN")
    chapter_two.content = "<h1>第二章</h1><p>第二章正文内容。</p>"
    book.add_item(chapter_one)
    book.add_item(chapter_two)
    book.toc = (chapter_one, chapter_two)
    book.spine = ["nav", chapter_one, chapter_two]
    book.add_item(epub.EpubNav())
    book.add_item(epub.EpubNcx())
    epub.write_epub(str(path), book)


def test_epub_spine_order_and_chapters(tmp_path: Path):
    source = tmp_path / "sample.epub"
    _create_epub(source)

    result = import_document(source)

    assert result.title == "电子书测试"
    assert [chapter.title for chapter in result.chapters][:2] == ["第一章", "第二章"]
    assert result.text.index("第一章正文") < result.text.index("第二章正文")


def test_real_mobi_conversion_and_chapter_order(tmp_path: Path):
    ebook_convert_candidates = [
        shutil.which("ebook-convert"),
        "/Applications/calibre.app/Contents/MacOS/ebook-convert",
    ]
    ebook_convert = next(
        (
            candidate
            for candidate in ebook_convert_candidates
            if candidate and Path(candidate).is_file()
        ),
        None,
    )
    if not ebook_convert:
        pytest.skip("Calibre ebook-convert is unavailable")

    source_epub = tmp_path / "mobi-source.epub"
    source_mobi = tmp_path / "sample.mobi"
    _create_epub(source_epub)
    process = subprocess.run(
        [ebook_convert, str(source_epub), str(source_mobi)],
        capture_output=True,
        text=True,
        timeout=300,
    )
    assert process.returncode == 0, process.stderr
    assert source_mobi.exists()

    result = import_document(source_mobi)

    assert result.file_type == "MOBI"
    assert "第一章正文内容" in result.text
    assert "第二章正文内容" in result.text
    assert result.text.index("第一章正文") < result.text.index("第二章正文")


def test_cache_round_trip(tmp_path: Path):
    source = tmp_path / "cached.txt"
    source.write_text("缓存往返测试文本。", encoding="utf-8")
    cache_dir = tmp_path / "cache"

    first = import_document(source, cache_dir=cache_dir)
    second = import_document(source, cache_dir=cache_dir)

    assert first.to_dict() == second.to_dict()
    assert len(list(cache_dir.glob("*.json"))) == 1


def test_rejects_unsupported_extension(tmp_path: Path):
    source = tmp_path / "unsafe.exe"
    source.write_bytes(b"not a document")

    with pytest.raises(DocumentImportError, match="不支持"):
        import_document(source)
