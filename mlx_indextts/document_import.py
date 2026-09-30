"""Offline document and ebook import for the IndexTTS2 WebUI."""

from __future__ import annotations

import hashlib
import html
import json
import re
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Callable


SUPPORTED_EXTENSIONS = {".txt", ".md", ".markdown", ".doc", ".docx", ".pdf", ".epub", ".mobi"}
MAX_FILE_BYTES = 500 * 1024 * 1024
MAX_ARCHIVE_UNCOMPRESSED_BYTES = 1024 * 1024 * 1024
MAX_ARCHIVE_MEMBERS = 20_000
CACHE_VERSION = 4


class DocumentImportError(RuntimeError):
    """Raised when an uploaded document cannot be safely parsed."""


@dataclass
class DocumentChapter:
    title: str
    text: str
    source_index: int
    page_start: int | None = None
    page_end: int | None = None
    include_title: bool = True

    @property
    def character_count(self) -> int:
        return count_effective_characters(self.text)


@dataclass
class ImportedDocument:
    filename: str
    file_type: str
    title: str
    chapters: list[DocumentChapter]
    used_ocr: bool = False
    ocr_pages: list[int] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def text(self) -> str:
        return join_chapters(self.chapters)

    @property
    def character_count(self) -> int:
        return count_effective_characters(self.text)

    @property
    def chinese_character_count(self) -> int:
        return count_chinese_characters(self.text)

    @property
    def word_count(self) -> int:
        return len(re.findall(r"\b[A-Za-z0-9]+(?:['-][A-Za-z0-9]+)*\b", self.text))

    def to_dict(self) -> dict:
        data = asdict(self)
        data["character_count"] = self.character_count
        data["chinese_character_count"] = self.chinese_character_count
        data["word_count"] = self.word_count
        return data

    @classmethod
    def from_dict(cls, data: dict) -> "ImportedDocument":
        return cls(
            filename=str(data["filename"]),
            file_type=str(data["file_type"]),
            title=str(data.get("title") or Path(str(data["filename"])).stem),
            chapters=[DocumentChapter(**chapter) for chapter in data.get("chapters", [])],
            used_ocr=bool(data.get("used_ocr", False)),
            ocr_pages=[int(page) for page in data.get("ocr_pages", [])],
            warnings=[str(warning) for warning in data.get("warnings", [])],
        )


def count_effective_characters(text: str) -> int:
    return sum(not character.isspace() for character in text)


def count_chinese_characters(text: str) -> int:
    return len(re.findall(r"[\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff]", text))


def estimate_audio_minutes(character_count: int, characters_per_minute: int = 330) -> float:
    return character_count / max(characters_per_minute, 1)


def join_chapters(chapters: list[DocumentChapter]) -> str:
    parts = []
    for chapter in chapters:
        text = clean_text(chapter.text)
        if not text:
            continue
        title = clean_text(chapter.title)
        if chapter.include_title and title and not text.startswith(title):
            parts.append(f"{title}\n\n{text}")
        else:
            parts.append(text)
    return "\n\n".join(parts).strip()


def clean_text(text: str) -> str:
    """Normalize extracted prose without collapsing paragraph structure."""
    text = html.unescape(text or "")
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\u00a0", " ")
    text = re.sub(r"[\u200b-\u200f\u202a-\u202e\u2060\ufeff]", "", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _safe_title(value: str, fallback: str) -> str:
    title = clean_text(value).split("\n", 1)[0][:120]
    return title or fallback


def _single_chapter(title: str, text: str) -> list[DocumentChapter]:
    cleaned = clean_text(text)
    if not cleaned:
        raise DocumentImportError("未从文件中读取到有效文字。")
    return [DocumentChapter(title=title, text=cleaned, source_index=0)]


def _validate_path(path: Path) -> str:
    if not path.exists() or not path.is_file():
        raise DocumentImportError("上传文件不存在或不是普通文件。")
    extension = path.suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise DocumentImportError(
            f"不支持 {extension or '未知'} 格式。可用格式："
            f"{', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )
    size = path.stat().st_size
    if size <= 0:
        raise DocumentImportError("文件为空。")
    if size > MAX_FILE_BYTES:
        raise DocumentImportError("文件超过 500 MB 安全限制。")
    return extension


def _validate_zip_archive(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > MAX_ARCHIVE_MEMBERS:
                raise DocumentImportError("压缩包内文件数过多，已停止解析。")
            total = sum(member.file_size for member in members)
            if total > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
                raise DocumentImportError("压缩包解压后超过 1 GB 安全限制。")
            for member in members:
                member_path = Path(member.filename)
                if member_path.is_absolute() or ".." in member_path.parts:
                    raise DocumentImportError("压缩包包含不安全路径。")
    except zipfile.BadZipFile as exc:
        raise DocumentImportError("文件已损坏或不是有效的压缩文档。") from exc


def _read_text_file(path: Path) -> str:
    from charset_normalizer import from_bytes

    raw = path.read_bytes()
    if raw.startswith(b"\xef\xbb\xbf"):
        return raw.decode("utf-8-sig")
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16")
    detected = from_bytes(raw).best()
    candidates: list[str] = []
    if detected is not None:
        candidates.append(str(detected))
    for encoding in ("utf-8", "gb18030", "big5", "shift_jis"):
        try:
            decoded = raw.decode(encoding)
        except UnicodeDecodeError:
            continue
        if decoded not in candidates:
            candidates.append(decoded)
    if not candidates:
        raise DocumentImportError("无法识别 TXT/MD 文件编码。")

    def readability_score(value: str) -> float:
        if not value:
            return float("-inf")
        chinese = count_chinese_characters(value)
        korean = len(re.findall(r"[\uac00-\ud7af]", value))
        japanese = len(re.findall(r"[\u3040-\u30ff]", value))
        controls = sum(ord(character) < 32 and character not in "\n\r\t" for character in value)
        replacements = value.count("\ufffd")
        printable = sum(character.isprintable() or character in "\n\r\t" for character in value)
        return (
            printable / len(value)
            + chinese * 2.0
            - korean * 1.5
            - japanese * 0.5
            - controls * 10.0
            - replacements * 20.0
        )

    return max(candidates, key=readability_score)


def _markdown_to_text(markdown: str) -> str:
    text = re.sub(r"```.*?```", "\n", markdown, flags=re.DOTALL)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"!\[([^\]]*)\]\([^)]*\)", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", text)
    text = re.sub(r"^\s{0,3}#{1,6}\s*", "", text, flags=re.MULTILINE)
    text = re.sub(r"^\s*>\s?", "", text, flags=re.MULTILINE)
    text = re.sub(r"^\s*(?:[-+*]|\d+[.)])\s+", "", text, flags=re.MULTILINE)
    text = re.sub(r"[*_~]{1,3}([^\n*_~]+)[*_~]{1,3}", r"\1", text)
    text = re.sub(r"^\s*[-*_]{3,}\s*$", "", text, flags=re.MULTILINE)
    text = re.sub(r"<[^>]+>", "", text)
    return clean_text(text)


def _parse_txt_or_markdown(path: Path, markdown: bool) -> ImportedDocument:
    raw = _read_text_file(path)
    text = _markdown_to_text(raw) if markdown else clean_text(raw)
    title = path.stem
    return ImportedDocument(
        filename=path.name,
        file_type="Markdown" if markdown else "TXT",
        title=title,
        chapters=_single_chapter(title, text),
    )


def _parse_docx(path: Path, display_filename: str | None = None) -> ImportedDocument:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    _validate_zip_archive(path)
    try:
        document = Document(str(path))
    except Exception as exc:
        raise DocumentImportError(f"Word 文档无法打开：{exc}") from exc

    chapters: list[DocumentChapter] = []
    current_title = path.stem
    current_parts: list[str] = []

    def flush() -> None:
        nonlocal current_parts
        body = clean_text("\n".join(current_parts))
        if body:
            chapters.append(
                DocumentChapter(current_title, body, source_index=len(chapters))
            )
        current_parts = []

    for element in document.iter_inner_content():
        if isinstance(element, Paragraph):
            text = clean_text(element.text)
            if not text:
                if current_parts and current_parts[-1] != "":
                    current_parts.append("")
                continue
            style_name = (element.style.name if element.style else "").lower()
            if style_name.startswith("heading") or style_name.startswith("标题"):
                flush()
                current_title = _safe_title(text, f"章节 {len(chapters) + 1}")
            else:
                current_parts.append(text)
        elif isinstance(element, Table):
            for row in element.rows:
                cells = [clean_text(cell.text) for cell in row.cells]
                if any(cells):
                    current_parts.append("，".join(cell for cell in cells if cell))
    flush()

    if not chapters:
        raise DocumentImportError("Word 文档中未发现可读取正文。")
    return ImportedDocument(
        filename=display_filename or path.name,
        file_type="Word DOCX",
        title=path.stem,
        chapters=chapters,
    )


def _find_soffice() -> str | None:
    candidates = [
        shutil.which("soffice"),
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
        "/Applications/LibreOfficeDev.app/Contents/MacOS/soffice",
    ]
    return next((candidate for candidate in candidates if candidate and Path(candidate).exists()), None)


def _parse_legacy_doc(path: Path) -> ImportedDocument:
    with tempfile.TemporaryDirectory(prefix="indextts_doc_") as temporary_directory:
        temp_dir = Path(temporary_directory)
        converted = temp_dir / f"{path.stem}.docx"
        soffice = _find_soffice()
        errors = []
        if soffice:
            process = subprocess.run(
                [
                    soffice,
                    "--headless",
                    "--convert-to",
                    "docx",
                    "--outdir",
                    str(temp_dir),
                    str(path),
                ],
                capture_output=True,
                text=True,
                timeout=120,
            )
            if process.returncode != 0:
                errors.append(process.stderr.strip() or process.stdout.strip())
        if converted.exists():
            result = _parse_docx(converted, display_filename=path.name)
            result.file_type = "Word DOC"
            result.title = path.stem
            return result

        converted_txt = temp_dir / f"{path.stem}.txt"
        process = subprocess.run(
            ["/usr/bin/textutil", "-convert", "txt", "-output", str(converted_txt), str(path)],
            capture_output=True,
            text=True,
            timeout=120,
        )
        if process.returncode == 0 and converted_txt.exists():
            text = _read_text_file(converted_txt)
            return ImportedDocument(
                filename=path.name,
                file_type="Word DOC",
                title=path.stem,
                chapters=_single_chapter(path.stem, text),
                warnings=["旧版 DOC 由系统文本转换器读取，复杂排版可能被简化。"],
            )
        errors.append(process.stderr.strip() or process.stdout.strip())
        detail = "；".join(error for error in errors if error)
        raise DocumentImportError(f"旧版 DOC 转换失败。{detail}")


def _remove_repeated_page_edges(page_texts: list[str]) -> list[str]:
    if len(page_texts) < 3:
        return page_texts
    edge_counts: dict[str, int] = {}
    split_pages = []
    for text in page_texts:
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        split_pages.append(lines)
        for line in lines[:2] + lines[-2:]:
            if 2 <= len(line) <= 100:
                edge_counts[line] = edge_counts.get(line, 0) + 1
    threshold = max(3, int(len(page_texts) * 0.6))
    repeated = {line for line, count in edge_counts.items() if count >= threshold}
    return [clean_text("\n".join(line for line in lines if line not in repeated)) for lines in split_pages]


def _ocr_pdf_page(page, page_number: int) -> str:
    import pymupdf

    tesseract = shutil.which("tesseract")
    if not tesseract:
        raise DocumentImportError("本机未找到 Tesseract OCR。")
    pixmap = page.get_pixmap(matrix=pymupdf.Matrix(2.5, 2.5), alpha=False)
    png_bytes = pixmap.tobytes("png")
    process = subprocess.run(
        [
            tesseract,
            "stdin",
            "stdout",
            "-l",
            "chi_sim+eng",
            "--psm",
            "6",
            "quiet",
        ],
        input=png_bytes,
        capture_output=True,
        timeout=180,
    )
    if process.returncode != 0:
        detail = process.stderr.decode("utf-8", errors="replace").strip()
        raise DocumentImportError(f"PDF 第 {page_number} 页 OCR 失败：{detail}")
    return clean_text(process.stdout.decode("utf-8", errors="replace"))


def _parse_pdf(
    path: Path,
    progress_callback: Callable[[int, int, str], None] | None = None,
) -> ImportedDocument:
    import pymupdf

    try:
        document = pymupdf.open(path)
    except Exception as exc:
        raise DocumentImportError(f"PDF 无法打开：{exc}") from exc
    if document.needs_pass:
        document.close()
        raise DocumentImportError("PDF 已加密，请先解除打开密码。")

    page_texts = []
    ocr_pages = []
    warnings = []
    total_pages = document.page_count
    for page_index, page in enumerate(document):
        if progress_callback:
            progress_callback(page_index, total_pages, f"正在读取 PDF 第 {page_index + 1}/{total_pages} 页")
        blocks = sorted(
            page.get_text("blocks", sort=True),
            key=lambda block: (round(block[1] / 12), block[0]),
        )
        text = clean_text("\n".join(str(block[4]) for block in blocks if len(block) > 4))
        if count_effective_characters(text) < 20:
            text = _ocr_pdf_page(page, page_index + 1)
            ocr_pages.append(page_index + 1)
        page_texts.append(text)
    document.close()
    page_texts = _remove_repeated_page_edges(page_texts)

    chapters = []
    for index, text in enumerate(page_texts):
        if not text:
            warnings.append(f"第 {index + 1} 页未读取到有效文字。")
            continue
        chapters.append(
            DocumentChapter(
                title=f"第 {index + 1} 页",
                text=text,
                source_index=index,
                page_start=index + 1,
                page_end=index + 1,
                include_title=False,
            )
        )
    if not chapters:
        raise DocumentImportError("PDF 文本提取和 OCR 均未读取到有效正文。")
    if ocr_pages:
        warnings.append(f"已对 {len(ocr_pages)} 页扫描内容执行本机 OCR。")
    return ImportedDocument(
        filename=path.name,
        file_type="PDF",
        title=path.stem,
        chapters=chapters,
        used_ocr=bool(ocr_pages),
        ocr_pages=ocr_pages,
        warnings=warnings,
    )


def _html_to_text(content: bytes | str) -> tuple[str, str]:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(content, "html.parser")
    for tag in soup(["script", "style", "nav", "noscript", "svg"]):
        tag.decompose()
    heading = soup.find(["h1", "h2", "title"])
    title = clean_text(heading.get_text(" ", strip=True)) if heading else ""
    text = clean_text(soup.get_text("\n", strip=True))
    return title, text


def _is_navigation_page(content: bytes) -> bool:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(content, "html.parser")
    for tag in soup(["script", "style", "noscript", "svg"]):
        tag.decompose()
    body = soup.body or soup
    links = body.find_all("a", href=True)
    if len(links) < 2:
        return False
    total = count_effective_characters(body.get_text(" ", strip=True))
    if total > 2000:
        return False
    link_lengths = [count_effective_characters(link.get_text(" ", strip=True)) for link in links]
    if any(length > 120 for length in link_lengths):
        return False
    linked = sum(link_lengths)
    return linked >= max(4, total // 4) and total - linked <= 40


def _parse_epub(path: Path, display_filename: str | None = None) -> ImportedDocument:
    import ebooklib
    from ebooklib import epub

    _validate_zip_archive(path)
    try:
        book = epub.read_epub(str(path), options={"ignore_ncx": True})
    except Exception as exc:
        raise DocumentImportError(f"EPUB 无法打开：{exc}") from exc

    metadata_title = book.get_metadata("DC", "title")
    book_title = clean_text(str(metadata_title[0][0])) if metadata_title else path.stem
    chapters = []
    seen_items = set()

    def add_item(item) -> None:
        if item is None or item.get_id() in seen_items:
            return
        seen_items.add(item.get_id())
        if item.get_type() != ebooklib.ITEM_DOCUMENT:
            return
        content = item.get_content()
        if content.count(b"href=") >= 2 and _is_navigation_page(content):
            return
        title, text = _html_to_text(content)
        if count_effective_characters(text) < 2:
            return
        if not title:
            first_line = text.split("\n", 1)[0].strip()
            if re.fullmatch(
                r"(?:第[一二三四五六七八九十百千万零〇两\d]{1,12}[章节回卷篇]|Chapter\s+\d+)"
                r"(?:[：:\s\-—·].{1,60})?",
                first_line,
                flags=re.IGNORECASE,
            ):
                title = first_line
        chapters.append(
            DocumentChapter(
                title=_safe_title(title, f"章节 {len(chapters) + 1}"),
                text=text,
                source_index=len(chapters),
            )
        )

    for item_id, _ in book.spine:
        add_item(book.get_item_with_id(item_id))
    for item in book.get_items_of_type(ebooklib.ITEM_DOCUMENT):
        add_item(item)
    if not chapters:
        raise DocumentImportError("EPUB 中未找到可读取的正文章节。")
    return ImportedDocument(
        filename=display_filename or path.name,
        file_type="EPUB",
        title=book_title or path.stem,
        chapters=chapters,
    )


def _parse_html_file(path: Path, display_filename: str) -> ImportedDocument:
    raw = path.read_bytes()
    title, text = _html_to_text(raw)
    final_title = _safe_title(title, Path(display_filename).stem)
    return ImportedDocument(
        filename=display_filename,
        file_type="MOBI",
        title=final_title,
        chapters=_single_chapter(final_title, text),
    )


def _find_ebook_convert() -> str | None:
    candidates = [
        shutil.which("ebook-convert"),
        "/Applications/calibre.app/Contents/MacOS/ebook-convert",
    ]
    return next(
        (candidate for candidate in candidates if candidate and Path(candidate).is_file()),
        None,
    )


def _parse_mobi(
    path: Path,
    progress_callback: Callable[[int, int, str], None] | None = None,
) -> ImportedDocument:
    if progress_callback:
        progress_callback(0, 1, "正在转换 MOBI 电子书")

    ebook_convert = _find_ebook_convert()
    if ebook_convert:
        with tempfile.TemporaryDirectory(prefix="indextts-mobi-") as temporary_directory:
            converted_epub = Path(temporary_directory) / "converted.epub"
            process = subprocess.run(
                [ebook_convert, str(path), str(converted_epub)],
                capture_output=True,
                text=True,
                timeout=300,
            )
            if process.returncode == 0 and converted_epub.exists():
                result = _parse_epub(converted_epub, display_filename=path.name)
                result.file_type = "MOBI"
                return result
            message = (process.stderr or process.stdout).strip()
            if "drm" in message.lower() or "encrypted" in message.lower():
                raise DocumentImportError("MOBI 已加密或含 DRM，无法读取。")

    import mobi

    temporary_directory = None
    try:
        temporary_directory, extracted_path = mobi.extract(str(path))
        extracted = Path(extracted_path)
        suffix = extracted.suffix.lower()
        if suffix == ".epub":
            result = _parse_epub(extracted, display_filename=path.name)
        elif suffix == ".pdf":
            result = _parse_pdf(extracted, progress_callback)
            result.filename = path.name
        elif suffix in {".html", ".htm", ".xhtml"}:
            result = _parse_html_file(extracted, path.name)
        else:
            candidates = list(Path(temporary_directory).rglob("*.epub"))
            candidates += list(Path(temporary_directory).rglob("*.html"))
            if not candidates:
                raise DocumentImportError("MOBI 解包后未找到 EPUB 或 HTML 正文。")
            candidate = candidates[0]
            result = (
                _parse_epub(candidate, display_filename=path.name)
                if candidate.suffix.lower() == ".epub"
                else _parse_html_file(candidate, path.name)
            )
        result.file_type = "MOBI"
        result.filename = path.name
        return result
    except DocumentImportError:
        raise
    except Exception as exc:
        message = str(exc)
        if "drm" in message.lower() or "encrypted" in message.lower():
            raise DocumentImportError("MOBI 已加密或含 DRM，无法读取。") from exc
        raise DocumentImportError(f"MOBI 解析失败：{message}") from exc
    finally:
        if temporary_directory:
            shutil.rmtree(temporary_directory, ignore_errors=True)


def _cache_key(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    digest.update(f"document-import-v{CACHE_VERSION}".encode())
    return digest.hexdigest()


def import_document(
    path: str | Path,
    *,
    cache_dir: str | Path | None = None,
    progress_callback: Callable[[int, int, str], None] | None = None,
) -> ImportedDocument:
    """Parse a supported document entirely on the local machine."""
    source = Path(path).resolve()
    extension = _validate_path(source)
    cache_path = None
    if cache_dir:
        cache_root = Path(cache_dir)
        cache_root.mkdir(parents=True, exist_ok=True)
        cache_path = cache_root / f"{_cache_key(source)}.json"
        if cache_path.exists():
            try:
                return ImportedDocument.from_dict(json.loads(cache_path.read_text("utf-8")))
            except (OSError, ValueError, TypeError, KeyError, json.JSONDecodeError):
                cache_path.unlink(missing_ok=True)

    if extension == ".txt":
        result = _parse_txt_or_markdown(source, markdown=False)
    elif extension in {".md", ".markdown"}:
        result = _parse_txt_or_markdown(source, markdown=True)
    elif extension == ".docx":
        result = _parse_docx(source)
    elif extension == ".doc":
        result = _parse_legacy_doc(source)
    elif extension == ".pdf":
        result = _parse_pdf(source, progress_callback)
    elif extension == ".epub":
        result = _parse_epub(source)
    elif extension == ".mobi":
        result = _parse_mobi(source, progress_callback)
    else:  # pragma: no cover - guarded by _validate_path
        raise DocumentImportError(f"未实现的格式：{extension}")

    _validate_extracted_document(result)
    if cache_path:
        temporary_path = cache_path.with_suffix(".tmp")
        temporary_path.write_text(
            json.dumps(result.to_dict(), ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        temporary_path.replace(cache_path)
    return result


def _validate_extracted_document(document: ImportedDocument) -> None:
    """Reject corrupt or obviously unreadable extraction before WebUI output."""
    text = document.text
    if document.character_count == 0:
        raise DocumentImportError("文档解析完成，但有效字数为 0。")
    replacement_ratio = text.count("\ufffd") / max(len(text), 1)
    if replacement_ratio > 0.001:
        raise DocumentImportError("文档包含过多乱码替换符，解析质量检查未通过。")
    mojibake_markers = ("锟斤拷", "锟斤拷", "ï¿½", "é", "æå­")
    if any(marker in text for marker in mojibake_markers):
        raise DocumentImportError("检测到典型乱码，解析质量检查未通过。")
    empty_chapters = sum(not clean_text(chapter.text) for chapter in document.chapters)
    if empty_chapters:
        raise DocumentImportError("存在空章节，解析质量检查未通过。")

    normalized_chapters = [
        re.sub(r"\s+", "", chapter.text) for chapter in document.chapters if chapter.text
    ]
    duplicate_count = len(normalized_chapters) - len(set(normalized_chapters))
    if duplicate_count and duplicate_count / max(len(normalized_chapters), 1) > 0.2:
        document.warnings.append(
            f"检测到 {duplicate_count} 个重复章节/页面，请在预览中复核。"
        )


def select_document_text(document: ImportedDocument | dict, selected_titles: list[str] | None) -> str:
    if isinstance(document, dict):
        document = ImportedDocument.from_dict(document)
    if not selected_titles:
        raise DocumentImportError("请至少选择一个章节。")
    selected_indices = set()
    choices = {
        f"{index + 1}. {chapter.title}": index
        for index, chapter in enumerate(document.chapters)
    }
    for title in selected_titles:
        if title in choices:
            selected_indices.add(choices[title])
            continue
        matches = [
            index for index, chapter in enumerate(document.chapters)
            if chapter.title == title
        ]
        if len(matches) != 1:
            raise DocumentImportError("章节选择无效或标题重复，请从章节列表选择具体章节。")
        selected_indices.add(matches[0])
    chapters = [
        chapter for index, chapter in enumerate(document.chapters)
        if index in selected_indices
    ]
    return join_chapters(chapters)
